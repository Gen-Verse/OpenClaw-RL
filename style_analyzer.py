import os
import re
from collections import Counter

def read_texts(directory):
    """Read text from common document types including .txt, .md, .docx, and .xlsx.
    Returns a list of file text contents.
    """
    texts = []
    for root, dirs, files in os.walk(directory):
        for fname in files:
            path = os.path.join(root, fname)
            lname = fname.lower()
            try:
                if lname.endswith(('.txt', '.md', '.markdown', '.rst', '.tex')):
                    with open(path, 'r', encoding='utf-8', errors='ignore') as fh:
                        texts.append(fh.read())
                elif lname.endswith('.docx'):
                    try:
                        from docx import Document
                        doc = Document(path)
                        para_text = '\n'.join(p.text for p in doc.paragraphs if p.text)
                        if para_text:
                            texts.append(para_text)
                    except Exception:
                        continue
                elif lname.endswith(('.xls', '.xlsx')):
                    try:
                        import openpyxl
                        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
                        rows = []
                        for ws in wb.worksheets:
                            for row in ws.iter_rows(values_only=True):
                                for cell in row:
                                    if cell is not None:
                                        rows.append(str(cell))
                        if rows:
                            texts.append('\n'.join(rows))
                    except Exception:
                        continue
            except Exception:
                # ignore unreadable files
                continue
    return texts

def analyze_directory(directory):
    texts = read_texts(directory)
    combined = '\n'.join(texts)
    words = re.findall(r"\w+", combined.lower())
    word_counts = Counter(words).most_common(50)
    sentences = re.split(r'[.!?]+', combined)
    sent_lens = [len(re.findall(r"\w+", s)) for s in sentences if s.strip()]
    avg_sent_len = (sum(sent_lens) / len(sent_lens)) if sent_lens else 0
    # bigrams
    bigrams = Counter(zip(words, words[1:]))
    top_bigrams = [' '.join(b) for b,_ in bigrams.most_common(10)]
    return {
        'num_files': len(texts),
        'num_words': len(words),
        'top_words': word_counts,
        'avg_sentence_length': avg_sent_len,
        'top_bigrams': top_bigrams
    }

def generate_sample(analysis, sentences=3):
    """生成示例文本。优先调用 OpenAI（若设置 OPENAI_API_KEY），否则使用简单的局部拼接回退。
    返回纯文本段落。
    """
    import os
    try:
        import openai
    except Exception:
        openai = None

    top_words = [w for w,_ in analysis.get('top_words', [])][:20]
    avg_len = int(analysis.get('avg_sentence_length', 12)) or 12
    if not top_words:
        return '（未检测到文本）'

    api_key = os.getenv('OPENAI_API_KEY')
    if api_key and openai:
        try:
            openai.api_key = api_key
            prompt = (
                f"你是一个写作风格仿写助手。根据以下写作风格特征生成{sentences}个句子，组成一段简洁、丝滑、自然的中文文本，避免复制原文内容：\n"
                f"常用词（前20）：{', '.join(top_words)}\n"
                f"平均句长（词）：{avg_len}\n"
                f"常见短语：{', '.join(analysis.get('top_bigrams', [])[:10])}\n"
                "要求：使用流畅自然的中文，语气简洁优雅，每句长度与给定平均句长相近，不要包含源文件具体句子或版权内容。只返回纯文本段落。"
            )
            model_name = os.getenv('OPENAI_MODEL', 'gpt-4')
            resp = openai.ChatCompletion.create(
                model=model_name,
                messages=[
                    {"role": "system", "content": "你是一个中文写作助手，擅长模仿风格。"},
                    {"role": "user", "content": prompt}
                ],
                max_tokens=max(150, avg_len * sentences * 8),
                temperature=0.8,
            )
            # Extract reply safely
            text = None
            if isinstance(resp, dict) and resp.get('choices'):
                choice = resp['choices'][0]
                if isinstance(choice.get('message'), dict):
                    text = choice['message'].get('content')
                else:
                    text = choice.get('text')
            if text:
                return text.strip()
        except Exception:
            # Fail quietly to fallback generator
            pass

    # Fallback simple generator (保留局部风格特征，但不使用 AI)
    out = []
    idx = 0
    for _ in range(sentences):
        sent = []
        for i in range(avg_len):
            sent.append(top_words[(idx + i) % len(top_words)])
        out.append((' '.join(sent)).capitalize() + '.')
        idx += max(1, avg_len // 2)
    return ' '.join(out)
